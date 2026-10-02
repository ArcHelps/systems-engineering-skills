"""Bounded local readers and source-preserving import/export operations."""

import csv
import hashlib
import json
import math
import os
import re
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5

from defusedxml import ElementTree
from defusedxml.common import DefusedXmlException
from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader


MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_RECORDS = 100_000
MAX_IMPORT_RECORDS = 10_000
SUPPORTED = {".csv", ".tsv", ".xlsx", ".reqif", ".xml", ".json", ".txt", ".md", ".pdf", ".docx"}


def scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def local_tag(element):
    return element.tag.rsplit("}", 1)[-1]


def xml_ref(element, name):
    if element is None:
        return ""
    return next((node.text or "" for node in element.iter() if local_tag(node) == name), "")


def xhtml_text(element):
    """Separate XHTML blocks without splitting inline fragments inside a word."""
    parts = []
    blocks = {"div", "p", "br", "li", "tr", "td", "th", "h1", "h2", "h3", "h4", "h5", "h6"}

    def walk(node):
        block = local_tag(node).lower() in blocks
        if block:
            parts.append(" ")
        parts.append(node.text or "")
        for child in node:
            walk(child)
            parts.append(child.tail or "")
        if block:
            parts.append(" ")

    walk(element)
    return " ".join("".join(parts).split())


def spreadsheet_cell(value, *, xlsx=False):
    """User text never becomes executable formulas in exported workbooks."""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, allow_nan=False)
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError(
            "Table values must be finite numbers; preserve unavailable values explicitly"
        )
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        value = "'" + value
    if xlsx and isinstance(value, str):
        if len(value) > 32767:
            raise ValueError("XLSX cell exceeds 32,767 characters; use JSON to preserve its text")
        if re.search(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", value):
            raise ValueError("XLSX cell contains unsupported control characters; use JSON")
    return value


def check_zip(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if (
            len(entries) > 10_000
            or sum(info.file_size for info in entries) > MAX_UNCOMPRESSED_BYTES
        ):
            raise ValueError("Archive exceeds 10,000 entries or 200 MiB expanded size")
        if any(info.flag_bits & 1 for info in entries):
            raise ValueError("Encrypted archives are unsupported")


class Workspace:
    def __init__(self, root: Path | str):
        root = Path(root).expanduser()
        if not root.is_dir():
            raise ValueError("Workspace must be an existing directory")
        self.root = root.resolve()

    def path(self, value: str, *, output: bool = False) -> Path:
        if not value or "\x00" in value:
            raise ValueError("A local file path is required")
        candidate = Path(value).expanduser()
        if not candidate.is_absolute():
            candidate = self.root / candidate
        resolved = candidate.resolve()
        if not resolved.is_relative_to(self.root) or resolved == self.root:
            raise ValueError("File path must stay inside the configured workspace")
        if output:
            # Reject symlinked output paths even when their targets are inside the workspace.
            current = candidate
            while current != self.root and current.is_relative_to(self.root):
                if current.is_symlink():
                    raise ValueError("Output paths cannot contain symlinks")
                current = current.parent
            if resolved.exists():
                raise FileExistsError("Output already exists; choose a new filename")
        else:
            if not resolved.is_file():
                raise ValueError("Input must be an existing regular file")
            if resolved.stat().st_size > MAX_FILE_BYTES:
                raise ValueError("Input exceeds the 50 MiB local reader limit; split it explicitly")
            if resolved.suffix.lower() not in SUPPORTED:
                raise ValueError(
                    f"Unsupported format {resolved.suffix}; supported: {sorted(SUPPORTED)}"
                )
        return resolved

    def _source(self, path):
        digest = hashlib.sha256()
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(chunk)
        return {
            "path": str(path.relative_to(self.root)),
            "sha256": digest.hexdigest(),
            "bytes": path.stat().st_size,
        }

    @staticmethod
    def _headers(values):
        headers = [str(value).strip() if value is not None else "" for value in values]
        if (
            not headers
            or any(not header for header in headers)
            or len(set(headers)) != len(headers)
        ):
            raise ValueError(
                "Headers must be nonempty and unique; correct the source header row first"
            )
        return headers

    def _tabular(self, path, sheet=""):
        warnings = []
        records = []
        if path.suffix.lower() in {".csv", ".tsv"}:
            with path.open(encoding="utf-8-sig", newline="") as stream:
                reader = csv.reader(
                    stream, delimiter="\t" if path.suffix.lower() == ".tsv" else ",", strict=True
                )
                headers = self._headers(next(reader, []))
                previous_line = reader.line_num
                for row in reader:
                    row_number = previous_line + 1
                    source = {"row": row_number}
                    if reader.line_num != row_number:
                        source["end_row"] = reader.line_num
                    previous_line = reader.line_num
                    if len(row) > len(headers):
                        raise ValueError(
                            f"Row {row_number} has extra cells; refusing silent data loss"
                        )
                    records.append(
                        {
                            "values": dict(zip(headers, row + [""] * (len(headers) - len(row)))),
                            "source": source,
                        }
                    )
                    if len(records) > MAX_RECORDS:
                        raise ValueError(
                            "Input exceeds 100,000 records; split the source explicitly"
                        )
            return {"headers": headers, "records": records, "warnings": warnings}
        check_zip(path)
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            names = workbook.sheetnames
            if sheet and sheet not in names:
                raise ValueError(f"Unknown sheet; available sheets: {names}")
            if not sheet and len(names) > 1:
                raise ValueError(f"Select a sheet explicitly; available sheets: {names}")
            sheet = sheet or names[0]
            # Producer dimension metadata can be stale; scan actual worksheet rows.
            workbook[sheet].reset_dimensions()
            cached[sheet].reset_dimensions()
            rows = workbook[sheet].iter_rows()
            cache_rows = cached[sheet].iter_rows(values_only=True)
            header_row = [cell.value for cell in next(rows, [])]
            while header_row and header_row[-1] is None:
                header_row.pop()
            headers = self._headers(header_row)
            next(cache_rows, None)
            missing_formula_caches = 0
            for row_number, (row, cache_row) in enumerate(zip(rows, cache_rows), 2):
                if any(cell.value is not None for cell in row[len(headers) :]):
                    raise ValueError(f"Row {row_number} contains cells with no header")
                values = {}
                formats = {}
                for key, cell in zip(headers, row):
                    value = cell.value
                    number_format = cell.number_format
                    if number_format and number_format != "General":
                        formats[key] = {"raw_value": scalar(value), "number_format": number_format}
                    if (
                        isinstance(value, (int, float))
                        and not isinstance(value, bool)
                        and re.fullmatch(r"0+", number_format or "")
                        and float(value).is_integer()
                    ):
                        value = str(int(value)).zfill(len(number_format))
                    values[key] = scalar(value)
                values.update({key: None for key in headers if key not in values})
                formulas = {
                    key: {"expression": cell.value, "cached_value": scalar(cache_row[index])}
                    for index, (key, cell) in enumerate(zip(headers, row))
                    if isinstance(cell.value, str) and cell.value.startswith("=")
                }
                if any(value["cached_value"] is None for value in formulas.values()):
                    missing_formula_caches += 1
                    if len(warnings) < 20:
                        warnings.append(
                            f"{sheet}!row {row_number}: formula has no cached value; not evaluated"
                        )
                records.append(
                    {
                        "values": values,
                        "source": {"sheet": sheet, "row": row_number},
                        **({"formulas": formulas} if formulas else {}),
                        **({"cell_formats": formats} if formats else {}),
                    }
                )
                if len(records) > MAX_RECORDS:
                    raise ValueError("Sheet exceeds 100,000 records")
            if missing_formula_caches > 20:
                warnings.append(
                    f"{missing_formula_caches} rows have formulas without cached values; "
                    "first 20 warning locations shown. Every row retains its formula metadata."
                )
            return {"headers": headers, "records": records, "warnings": warnings, "sheet": sheet}
        finally:
            workbook.close()
            cached.close()

    def _reqif(self, path):
        try:
            root = ElementTree.parse(path).getroot()
        except (DefusedXmlException, ElementTree.ParseError) as error:
            raise ValueError(f"Unsafe or malformed ReqIF XML: {error}") from error
        if local_tag(root) != "REQ-IF":
            raise ValueError(
                "XML input must be ReqIF; arbitrary XML has no implicit engineering mapping"
            )
        bounded_tags = {
            "SPEC-RELATION",
            "SPEC-HIERARCHY",
            "SPEC-OBJECT-TYPE",
            "SPEC-RELATION-TYPE",
            "SPECIFICATION-TYPE",
            "ENUM-VALUE",
        }
        counts = {}
        for node in root.iter():
            tag = local_tag(node)
            if tag in bounded_tags or (
                tag.startswith("DATATYPE-DEFINITION-") and not tag.endswith("-REF")
            ):
                counts[tag] = counts.get(tag, 0) + 1
                if counts[tag] > MAX_RECORDS:
                    raise ValueError(
                        f"ReqIF exceeds {MAX_RECORDS:,} {tag} records; split it explicitly"
                    )
        definitions = {
            element.attrib.get("IDENTIFIER"): element.attrib.get("LONG-NAME")
            or element.attrib.get("IDENTIFIER")
            for element in root.iter()
            if local_tag(element).startswith("ATTRIBUTE-DEFINITION-")
            and not local_tag(element).endswith("-REF")
        }
        enums = {
            element.attrib.get("IDENTIFIER"): element.attrib.get("LONG-NAME")
            or element.attrib.get("IDENTIFIER")
            for element in root.iter()
            if local_tag(element) == "ENUM-VALUE"
        }
        objects = []
        headers = ["IDENTIFIER", "LONG-NAME"]
        warnings = []
        for element in root.iter():
            if local_tag(element) != "SPEC-OBJECT":
                continue
            values = {
                "IDENTIFIER": element.attrib.get("IDENTIFIER", ""),
                "LONG-NAME": element.attrib.get("LONG-NAME", ""),
            }
            raw = {}
            for attribute in element.iter():
                tag = local_tag(attribute)
                if not tag.startswith("ATTRIBUTE-VALUE-"):
                    continue
                definition = next(
                    (
                        node.text or ""
                        for node in attribute.iter()
                        if local_tag(node).startswith("ATTRIBUTE-DEFINITION-")
                        and local_tag(node).endswith("-REF")
                    ),
                    "",
                )
                key = definitions.get(definition, definition)
                if not key or key in values:
                    raise ValueError(
                        "ReqIF has missing or duplicate attribute names; map a corrected source"
                    )
                if tag == "ATTRIBUTE-VALUE-XHTML":
                    value_node = next(
                        (node for node in attribute if local_tag(node) == "THE-VALUE"), None
                    )
                    value = "" if value_node is None else xhtml_text(value_node)
                    raw[key] = (
                        ""
                        if value_node is None
                        else ElementTree.tostring(value_node, encoding="unicode")
                    )
                elif tag == "ATTRIBUTE-VALUE-ENUMERATION":
                    value = [
                        enums.get(node.text, node.text)
                        for node in attribute.iter()
                        if local_tag(node) == "ENUM-VALUE-REF"
                    ]
                else:
                    value = attribute.attrib.get("THE-VALUE", "")
                values[key] = value
                if key not in headers:
                    headers.append(key)
            objects.append(
                {
                    "values": values,
                    "source": {"object_id": values["IDENTIFIER"]},
                    "raw_xhtml": raw,
                    "object_type": xml_ref(element, "SPEC-OBJECT-TYPE-REF"),
                }
            )
            if len(objects) > MAX_RECORDS:
                raise ValueError("ReqIF exceeds 100,000 objects")
        relations = [
            {
                "id": node.attrib.get("IDENTIFIER", ""),
                "source": xml_ref(
                    next((child for child in node if local_tag(child) == "SOURCE"), None),
                    "SPEC-OBJECT-REF",
                ),
                "target": xml_ref(
                    next((child for child in node if local_tag(child) == "TARGET"), None),
                    "SPEC-OBJECT-REF",
                ),
                "type": xml_ref(node, "SPEC-RELATION-TYPE-REF"),
                "raw_xml": ElementTree.tostring(node, encoding="unicode"),
            }
            for node in root.iter()
            if local_tag(node) == "SPEC-RELATION"
        ]
        hierarchy = []

        def walk(node, ancestors, specification):
            if local_tag(node) == "SPECIFICATION":
                specification = node.attrib.get("IDENTIFIER", "")
            if local_tag(node) == "SPEC-HIERARCHY":
                object_node = next((child for child in node if local_tag(child) == "OBJECT"), None)
                object_id = (
                    xml_ref(object_node, "SPEC-OBJECT-REF") if object_node is not None else ""
                )
                hierarchy.append(
                    {
                        "id": node.attrib.get("IDENTIFIER", ""),
                        "object_id": object_id,
                        "parents": ancestors,
                        "specification": specification,
                    }
                )
                ancestors = ancestors + [object_id] if object_id else ancestors
            for child in node:
                walk(child, ancestors, specification)

        walk(root, [], "")
        if any(not row["source"] or not row["target"] for row in relations):
            warnings.append(
                "ReqIF contains relations with missing endpoints; retained for resolution"
            )
        if relations or hierarchy:
            warnings.append(
                "ReqIF relation/hierarchy semantics retained; they are not automatically derives links"
            )
        if any(local_tag(node).lower() in {"embedded-value", "object"} for node in root.iter()):
            warnings.append("Embedded attachments are not extracted; inspect the original source")
        type_definitions = [
            {
                "id": node.attrib.get("IDENTIFIER", ""),
                "name": node.attrib.get("LONG-NAME", ""),
                "kind": local_tag(node),
                "raw_xml": ElementTree.tostring(node, encoding="unicode"),
            }
            for node in root.iter()
            if local_tag(node) in {"SPEC-OBJECT-TYPE", "SPEC-RELATION-TYPE", "SPECIFICATION-TYPE"}
        ]
        datatype_definitions = [
            {
                "id": node.attrib.get("IDENTIFIER", ""),
                "name": node.attrib.get("LONG-NAME", ""),
                "kind": local_tag(node),
                "raw_xml": ElementTree.tostring(node, encoding="unicode"),
            }
            for node in root.iter()
            if local_tag(node).startswith("DATATYPE-DEFINITION-")
            and not local_tag(node).endswith("-REF")
        ]
        return {
            "headers": headers,
            "records": objects,
            "relations": relations,
            "hierarchy": hierarchy,
            "type_definitions": type_definitions,
            "datatype_definitions": datatype_definitions,
            "warnings": warnings,
        }

    def _read(self, path, sheet=""):
        suffix = path.suffix.lower()
        if suffix in {".csv", ".tsv", ".xlsx"}:
            try:
                return self._tabular(path, sheet)
            except csv.Error as error:
                raise ValueError(f"Malformed delimited file: {error}") from error
        if suffix in {".reqif", ".xml"}:
            return self._reqif(path)
        if suffix == ".json":

            def reject_constant(value):
                raise ValueError(f"Nonfinite JSON value {value}")

            def unique_keys(pairs):
                result = {}
                for key, value in pairs:
                    if key in result:
                        raise ValueError(f"Duplicate JSON key: {key}")
                    result[key] = value
                return result

            data = json.loads(
                path.read_text(encoding="utf-8-sig"),
                parse_constant=reject_constant,
                object_pairs_hook=unique_keys,
            )
            model_data = data.get("model", data) if isinstance(data, dict) else data
            rows = (
                model_data
                if isinstance(model_data, list)
                else model_data.get("items", [])
                if isinstance(model_data, dict)
                else []
            )
            if isinstance(rows, list) and rows and all(isinstance(row, dict) for row in rows):
                if len(rows) > MAX_RECORDS:
                    raise ValueError("JSON exceeds 100,000 records")
                headers = list(dict.fromkeys(key for row in rows for key in row))
                return {
                    "headers": headers,
                    "records": [
                        {"values": row, "source": {"index": index}}
                        for index, row in enumerate(rows)
                    ],
                    "warnings": [],
                    "document": data,
                    "relations": model_data.get("relationships", [])
                    if isinstance(model_data, dict)
                    else [],
                }
            return {"text": json.dumps(data, ensure_ascii=False, indent=2), "warnings": []}
        if suffix == ".pdf":
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise ValueError("Encrypted PDFs are unsupported")
            pages = [page.extract_text() or "" for page in reader.pages]
            warnings = [
                f"Page {index + 1}: no extractable text; OCR is required"
                for index, text in enumerate(pages)
                if not text.strip()
            ]
            return {
                "text": "\n".join(
                    f"[Page {index + 1}]\n{text}" for index, text in enumerate(pages)
                ),
                "pages": len(pages),
                "warnings": warnings,
            }
        if suffix == ".docx":
            check_zip(path)
            document = Document(path)
            # Paragraph/table order is preserved through the document body iterator.
            blocks = []
            for block in document.iter_inner_content():
                if hasattr(block, "text"):
                    blocks.append(block.text)
                else:
                    blocks.extend("\t".join(cell.text for cell in row.cells) for row in block.rows)
            return {
                "text": "\n".join(blocks),
                "warnings": [
                    "DOCX extraction covers body paragraphs/tables; drawings, comments, headers and tracked revisions require original-document inspection"
                ],
            }
        return {"text": path.read_text(encoding="utf-8-sig"), "warnings": []}

    def inspect_file(self, path: str) -> dict:
        resolved = self.path(path)
        result = {"source": self._source(resolved), "format": resolved.suffix.lower()}
        if resolved.suffix.lower() == ".xlsx":
            check_zip(resolved)
            workbook = load_workbook(resolved, read_only=True, keep_links=False)
            try:
                result["sheets"] = []
                for sheet in workbook:
                    sheet.reset_dimensions()
                    row_count, column_count, headers = 0, 0, []
                    for row_count, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if row_count > MAX_RECORDS + 1:
                            raise ValueError(
                                "Sheet exceeds 100,000 rows; split the source explicitly"
                            )
                        column_count = max(column_count, len(row))
                        if row_count == 1:
                            headers = [scalar(value) for value in row]
                    result["sheets"].append(
                        {
                            "name": sheet.title,
                            "rows": row_count,
                            "columns": column_count,
                            "headers": headers,
                        }
                    )
            finally:
                workbook.close()
            return result
        data = self._read(resolved)
        result.update(
            {key: value for key, value in data.items() if key in {"headers", "warnings", "pages"}}
        )
        if "records" in data:
            result["total_records"] = len(data["records"])
            result["relation_count"] = len(data.get("relations", []))
        else:
            result["characters"] = len(data.get("text", ""))
        return result

    def read_file(self, path: str, sheet: str = "", offset: int = 0, limit: int = 100) -> dict:
        if offset < 0 or not 1 <= limit <= 1000:
            raise ValueError("offset must be nonnegative; limit must be 1..1000")
        resolved = self.path(path)
        data = self._read(resolved, sheet)
        result = {"source": self._source(resolved), "warnings": data.get("warnings", [])}
        if "records" in data:
            rows = data["records"]
            result.update(
                {
                    "headers": data["headers"],
                    "records": rows[offset : offset + limit],
                    "total_records": len(rows),
                    "next_offset": offset + limit if offset + limit < len(rows) else None,
                }
            )
            for key in ("relations", "hierarchy", "type_definitions", "datatype_definitions"):
                if key in data:
                    result[key] = data[key][offset : offset + limit]
                    result[f"total_{key}"] = len(data[key])
                    result[f"next_{key}_offset"] = (
                        offset + limit if offset + limit < len(data[key]) else None
                    )
        else:
            text = data.get("text", "")
            # Text offsets and limits are in lines, not bytes or characters.
            lines = text.splitlines()
            result.update(
                {
                    "text": "\n".join(lines[offset : offset + limit]),
                    "offset_unit": "lines",
                    "total_lines": len(lines),
                    "next_offset": offset + limit if offset + limit < len(lines) else None,
                }
            )
        return result

    def import_requirements(self, path: str, mapping: dict[str, str], sheet: str = "") -> dict:
        allowed = {
            "id",
            "name",
            "statement",
            "parent",
            "notes",
            "rationale",
            "priority",
            "requirement_types",
        }
        if not isinstance(mapping, dict) or not mapping.get("statement") or set(mapping) - allowed:
            raise ValueError(
                "mapping requires statement and may include id/name/parent/notes/rationale/priority/requirement_types"
            )
        if any(not isinstance(column, str) for column in mapping.values()):
            raise ValueError("Mapping column names must be strings")
        resolved = self.path(path)
        source = self._source(resolved)
        data = self._read(resolved, sheet)
        if "records" not in data:
            raise ValueError(
                "Import requires tabular records; extract document requirements with the assistant first"
            )
        if set(mapping.values()) - set(data["headers"]):
            raise ValueError(f"Unknown mapped columns; available: {data['headers']}")
        if len(data["records"]) > MAX_IMPORT_RECORDS:
            raise ValueError(
                "Import exceeds 10,000 records; split source explicitly (nothing was imported)"
            )
        items, empty, rejected, warnings = [], 0, [], list(data.get("warnings", []))
        seen_local_ids = set()
        references = {}
        pending = []
        valid_types = {
            "Functional",
            "Performance",
            "Interface",
            "Design Criteria",
            "Environmental",
            "Operational",
            "Safety",
            "Compliance",
            "Quality",
            "Reliability",
            "Security",
            "Physical",
            "Human Factors",
            "Maintainability",
            "Mission",
            "Stakeholder",
            "Software",
            "Hardware",
            "Derived",
        }
        for index, row in enumerate(data["records"]):
            values = row["values"]
            if all(value is None or value == "" for value in values.values()):
                empty += 1
                continue
            statement = values.get(mapping["statement"], "")
            if not isinstance(statement, str) or not statement.strip():
                rejected.append(
                    {
                        "source": row["source"],
                        "values": values,
                        "reason": "Missing or non-text requirement statement",
                    }
                )
                continue
            source_id = values.get(mapping.get("id", ""))
            source_id = str(source_id) if source_id is not None and source_id != "" else None
            identity = source_id if source_id else f"row:{index}"
            local_id = "local:" + str(uuid5(NAMESPACE_URL, f"{source['path']}:{sheet}:{identity}"))
            if local_id in seen_local_ids:
                local_id = "local:" + str(
                    uuid5(NAMESPACE_URL, f"{source['path']}:{sheet}:{identity}:row:{index}")
                )
            seen_local_ids.add(local_id)
            name = values.get(mapping.get("name", "")) or source_id or statement[:80]
            fields = {"statement": statement}
            for field in ("notes", "rationale"):
                if field in mapping and values.get(mapping[field]) not in (None, ""):
                    fields[field] = str(values[mapping[field]])
            priority = values.get(mapping.get("priority", ""))
            if isinstance(priority, str) and priority in {"low", "medium", "high"}:
                fields["priority"] = priority
            elif priority not in (None, ""):
                warnings.append(
                    f"{source_id or index}: unmapped priority {priority!r}; retained in source fields"
                )
            kinds = values.get(mapping.get("requirement_types", ""))
            if kinds:
                tokens = (
                    kinds
                    if isinstance(kinds, list)
                    else [part.strip() for part in str(kinds).split(";")]
                )
                if set(tokens) <= valid_types:
                    fields["requirement_types"] = list(dict.fromkeys(tokens))
                else:
                    warnings.append(
                        f"{source_id or index}: unknown requirement types retained for review"
                    )
            item = {
                "id": local_id,
                "import_id": source_id,
                "type": "requirement",
                "name": str(name),
                "status": "draft",
                "fields": fields,
                "metadata": {
                    "source": {**source, **row["source"]},
                    "original_statement": statement,
                    "source_fields": values,
                    "imported_fields": {
                        key: value for key, value in values.items() if key not in mapping.values()
                    },
                },
            }
            for key in ("raw_xhtml", "formulas", "cell_formats", "object_type"):
                if key in row:
                    item["metadata"][key] = row[key]
            items.append(item)
            if source_id:
                references.setdefault(source_id, []).append(local_id)
            else:
                warnings.append(
                    f"Record {index}: missing source ID; local proposal identity supplied"
                )
            parent = values.get(mapping.get("parent", ""))
            if parent not in (None, ""):
                pending.append((local_id, str(parent)))
        for source_id, matches in references.items():
            if len(matches) > 1:
                warnings.append(
                    f"Duplicate source ID {source_id!r}; records preserved, links require resolution"
                )
        for item in items:
            if not item["import_id"] or len(references[item["import_id"]]) > 1:
                item["metadata"]["identity_unresolved"] = True
        relationships, unresolved = [], []
        for child, parent in pending:
            matches = references.get(parent, [])
            if len(matches) != 1 or child == (matches[0] if matches else ""):
                unresolved.append(
                    {
                        "child": child,
                        "parent_import_id": parent,
                        "reason": "Missing, duplicate, or self parent",
                    }
                )
                continue
            relationships.append(
                {
                    "id": "local:" + str(uuid5(NAMESPACE_URL, f"{matches[0]}:{child}:derives")),
                    "type": "derives",
                    "source": matches[0],
                    "target": child,
                    "fields": {},
                }
            )
        model = {"schema_version": "1.0", "items": items, "relationships": relationships}
        from .model import validate_model

        validation = validate_model(model)
        return {
            "model": model,
            "counts": {
                "source_records": len(data["records"]),
                "imported": len(items),
                "empty_records": empty,
                "rejected": len(rejected),
            },
            "rejected_records": rejected,
            "unresolved_links": unresolved,
            "warnings": warnings,
            "retained_reqif_relations": data.get("relations", []),
            "retained_reqif_hierarchy": data.get("hierarchy", []),
            "validation": validation,
            "proposal_only": True,
        }

    def export_artifact(self, path: str, data: dict | list | str) -> dict:
        target = self.path(path, output=True)
        if target.suffix.lower() not in {".json", ".csv", ".xlsx", ".md", ".txt"}:
            raise ValueError("Export supports JSON, CSV, XLSX, Markdown, or text")
        target.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary = tempfile.mkstemp(prefix=".arc-export-", dir=target.parent)
        os.close(descriptor)
        temporary = Path(temporary)
        try:
            suffix = target.suffix.lower()
            if suffix == ".json":
                temporary.write_text(
                    json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False),
                    encoding="utf-8",
                )
            elif suffix in {".md", ".txt"}:
                if not isinstance(data, str):
                    raise ValueError("Markdown/text export requires a string")
                temporary.write_text(data, encoding="utf-8")
            else:
                sheets = {"Results": data} if isinstance(data, list) else data
                if not isinstance(sheets, dict) or not sheets:
                    raise ValueError(
                        "Table export requires rows or a mapping of sheet names to rows"
                    )
                if suffix == ".csv" and len(sheets) != 1:
                    raise ValueError("CSV supports one table; use XLSX for multiple sheets")
                workbook = Workbook()
                workbook.remove(workbook.active)
                for name, rows in sheets.items():
                    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
                        raise ValueError("Each table must be a list of row objects")
                    headers = list(dict.fromkeys(key for row in rows for key in row))
                    if not headers:
                        raise ValueError("Table has no columns; use JSON for an empty result")
                    if not all(isinstance(header, str) for header in headers):
                        raise ValueError("Table column names must be strings")
                    if suffix == ".csv":
                        with temporary.open("w", encoding="utf-8", newline="") as stream:
                            writer = csv.DictWriter(stream, fieldnames=headers)
                            writer.writerow(
                                {header: spreadsheet_cell(header) for header in headers}
                            )
                            writer.writerows(
                                {key: spreadsheet_cell(value) for key, value in row.items()}
                                for row in rows
                            )
                    else:
                        if (
                            not isinstance(name, str)
                            or not name
                            or len(name) > 31
                            or any(c in name for c in "[]:*?/\\")
                            or name.lower() in {title.lower() for title in workbook.sheetnames}
                        ):
                            raise ValueError(
                                "XLSX sheet names must be nonempty, unique ignoring case, valid, and at most 31 characters"
                            )
                        sheet = workbook.create_sheet(name)
                        sheet.append([spreadsheet_cell(header, xlsx=True) for header in headers])
                        for row in rows:
                            sheet.append(
                                [spreadsheet_cell(row.get(key), xlsx=True) for key in headers]
                            )
                        sheet.freeze_panes = "A2"
                        sheet.auto_filter.ref = sheet.dimensions
                if suffix == ".xlsx":
                    workbook.save(temporary)
                workbook.close()
            # Atomic publish without overwrite, even if another writer wins the same name.
            os.link(temporary, target)
            return {
                "path": str(target.relative_to(self.root)),
                "source": self._source(target),
                "created": True,
            }
        finally:
            temporary.unlink(missing_ok=True)
