"""Behavioral acceptance checks, written before the implementation."""

import copy
import csv
import hashlib
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from arc_engineering.catalog import Catalog
from arc_engineering.files import Workspace
from arc_engineering.model import compare_models, trace_relationships, validate_model


def sample_model():
    return {
        "schema_version": "1.0",
        "items": [
            {"id": "P", "type": "requirement", "name": "Parent", "fields": {"statement": "Limit"}},
            {"id": "C", "type": "requirement", "name": "Child", "fields": {"statement": "Detail"}},
            {"id": "S", "type": "system", "name": "Sensor", "fields": {}},
        ],
        "relationships": [
            {"id": "D", "type": "derives", "source": "P", "target": "C", "fields": {}},
            {"id": "A", "type": "allocated_to", "source": "C", "target": "S", "fields": {}},
        ],
    }


class FileTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.workspace = Workspace(self.root)

    def tearDown(self):
        self.temp.cleanup()

    def test_workspace_rejects_outside_paths_symlinks_and_existing_outputs(self):
        with tempfile.TemporaryDirectory() as other:
            outside = Path(other) / "secret.txt"
            outside.write_text("private")
            (self.root / "link.txt").symlink_to(outside)
            for path in [str(outside), "../secret.txt", "link.txt"]:
                with self.assertRaises(ValueError):
                    self.workspace.read_file(path)
        self.workspace.export_artifact("result.json", {"ok": True})
        with self.assertRaises(FileExistsError):
            self.workspace.export_artifact("result.json", {"ok": False})
        self.assertEqual(json.loads((self.root / "result.json").read_text()), {"ok": True})

    def test_csv_import_preserves_originals_counts_unmapped_values_and_unresolved_parents(self):
        path = self.root / "req.csv"
        with path.open("w", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["ID", "Text", "Parent", "Vendor"])
            writer.writerows(
                [
                    ["001", "The sensor shall respond.", "MISSING", "Acme"],
                    ["001", "Second obligation", "", "Beta"],
                    ["", "Third obligation", "", "Gamma"],
                    ["", "", "", ""],
                ]
            )
        result = self.workspace.import_requirements(
            "req.csv", {"id": "ID", "statement": "Text", "parent": "Parent"}
        )
        self.assertEqual(len(result["model"]["items"]), 3)
        self.assertEqual(result["counts"]["source_records"], 4)
        self.assertEqual(result["counts"]["empty_records"], 1)
        first = result["model"]["items"][0]
        self.assertEqual(first["import_id"], "001")
        self.assertEqual(first["fields"]["statement"], "The sensor shall respond.")
        self.assertEqual(first["metadata"]["imported_fields"]["Vendor"], "Acme")
        self.assertTrue(result["unresolved_links"])
        self.assertTrue(any("duplicate" in warning.lower() for warning in result["warnings"]))
        self.assertEqual(validate_model(result["model"])["valid"], True)

    def test_duplicate_headers_and_bad_mappings_are_rejected_without_silent_loss(self):
        (self.root / "bad.csv").write_text("ID,Text,Text\n1,a,b\n")
        with self.assertRaises(ValueError):
            self.workspace.read_file("bad.csv")
        (self.root / "good.csv").write_text("ID,Text\n1,a\n")
        with self.assertRaises(ValueError):
            self.workspace.import_requirements("good.csv", {"statement": "Absent"})

    def test_workbook_formula_is_preserved_and_missing_cache_is_reported(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Requirements"
        sheet.append(["ID", "Text", "Limit"])
        sheet.append(["001", "Respond within budget", "=1+2"])
        workbook.save(self.root / "input.xlsx")
        result = self.workspace.read_file("input.xlsx", sheet="Requirements")
        self.assertEqual(result["records"][0]["values"]["Limit"], "=1+2")
        self.assertTrue(result["warnings"])
        self.assertEqual(result["records"][0]["source"]["row"], 2)

    def test_reqif_resolves_attribute_definitions_hierarchy_and_relation(self):
        xml = """<REQ-IF xmlns="http://www.omg.org/spec/ReqIF/20110401/reqif.xsd">
        <CORE-CONTENT><REQ-IF-CONTENT><SPEC-TYPES><SPEC-OBJECT-TYPE IDENTIFIER="T">
        <SPEC-ATTRIBUTES><ATTRIBUTE-DEFINITION-XHTML IDENTIFIER="AD" LONG-NAME="Text"/>
        </SPEC-ATTRIBUTES></SPEC-OBJECT-TYPE></SPEC-TYPES><SPEC-OBJECTS>
        <SPEC-OBJECT IDENTIFIER="P" LONG-NAME="Parent"><VALUES><ATTRIBUTE-VALUE-XHTML>
        <DEFINITION><ATTRIBUTE-DEFINITION-XHTML-REF>AD</ATTRIBUTE-DEFINITION-XHTML-REF>
        </DEFINITION><THE-VALUE><div xmlns="http://www.w3.org/1999/xhtml">Parent <b>text</b>
        </div></THE-VALUE></ATTRIBUTE-VALUE-XHTML></VALUES></SPEC-OBJECT>
        <SPEC-OBJECT IDENTIFIER="C" LONG-NAME="Child"/></SPEC-OBJECTS>
        <SPEC-RELATIONS><SPEC-RELATION IDENTIFIER="R"><SOURCE><SPEC-OBJECT-REF>P</SPEC-OBJECT-REF>
        </SOURCE><TARGET><SPEC-OBJECT-REF>C</SPEC-OBJECT-REF></TARGET></SPEC-RELATION>
        </SPEC-RELATIONS></REQ-IF-CONTENT></CORE-CONTENT></REQ-IF>"""
        (self.root / "input.reqif").write_text(xml)
        result = self.workspace.read_file("input.reqif")
        self.assertIn("Parent", result["records"][0]["values"]["Text"])
        self.assertEqual(result["relations"][0]["target"], "C")
        self.assertIn("R", result["relations"][0]["id"])

    def test_export_prevents_spreadsheet_formula_injection_and_leaves_source_unchanged(self):
        data = [{"text": '=WEBSERVICE("http://example.test")', "limit": -2}]
        self.workspace.export_artifact("out.csv", data)
        with (self.root / "out.csv").open(newline="") as stream:
            row = next(csv.DictReader(stream))
        self.assertTrue(row["text"].startswith("'="))
        self.workspace.export_artifact("out.xlsx", data)
        workbook = load_workbook(self.root / "out.xlsx", data_only=False)
        self.assertNotEqual(workbook.active["A2"].data_type, "f")
        workbook.close()
        self.assertEqual(data[0]["limit"], -2)

    def test_pages_are_explicit_and_xml_entities_are_rejected(self):
        (self.root / "rows.csv").write_text("ID,Text\n1,a\n2,b\n3,c\n")
        result = self.workspace.read_file("rows.csv", offset=1, limit=1)
        self.assertEqual(result["total_records"], 3)
        self.assertEqual(result["next_offset"], 2)
        self.assertEqual(result["records"][0]["values"]["ID"], "2")
        (self.root / "bad.reqif").write_text(
            '<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///etc/passwd">]><REQ-IF>&secret;</REQ-IF>'
        )
        with self.assertRaises(ValueError):
            self.workspace.read_file("bad.reqif")

    def test_multiline_csv_preserves_physical_source_lines(self):
        (self.root / "multiline.csv").write_text('ID,Text\nR1,"First\nsecond"\nR2,Third\n')
        result = self.workspace.read_file("multiline.csv")
        self.assertEqual(result["records"][0]["source"], {"row": 2, "end_row": 3})
        self.assertEqual(result["records"][1]["source"], {"row": 4})

    def test_malformed_csv_quotes_reject_instead_of_merging_requirements(self):
        (self.root / "bad.csv").write_text('ID,Text\nR1,"The unit shall stop\nR2,Second\n')
        with self.assertRaises(ValueError):
            self.workspace.read_file("bad.csv")

    def test_ambiguous_import_ids_cannot_be_compared_as_stable_identities(self):
        source = self.root / "duplicates.csv"
        source.write_text("ID,Text\nR1,First\nR1,Second\n")
        before = self.workspace.import_requirements(
            "duplicates.csv", {"id": "ID", "statement": "Text"}
        )
        source.write_text("ID,Text\nR1,Second\nR1,First\n")
        after = self.workspace.import_requirements(
            "duplicates.csv", {"id": "ID", "statement": "Text"}
        )
        self.assertTrue(
            all(item["metadata"]["identity_unresolved"] for item in before["model"]["items"])
        )
        with self.assertRaises(ValueError):
            compare_models(before["model"], after["model"])

    def test_table_export_rejects_nonfinite_numbers_without_publishing(self):
        for suffix in ("xlsx", "csv"):
            with self.subTest(suffix=suffix), self.assertRaises(ValueError):
                self.workspace.export_artifact("bad." + suffix, [{"value": float("nan")}])
            self.assertFalse((self.root / ("bad." + suffix)).exists())

    def test_reqif_xhtml_separates_blocks_and_preserves_inline_word_fragments(self):
        xml = """<REQ-IF><SPEC-OBJECT IDENTIFIER="R1"><VALUES><ATTRIBUTE-VALUE-XHTML>
        <DEFINITION><ATTRIBUTE-DEFINITION-XHTML-REF>AD</ATTRIBUTE-DEFINITION-XHTML-REF>
        </DEFINITION><THE-VALUE><div><p>The unit shall stop</p><p>and report a
        fault.</p><p>shut<b>down</b></p></div></THE-VALUE></ATTRIBUTE-VALUE-XHTML>
        </VALUES></SPEC-OBJECT></REQ-IF>"""
        (self.root / "blocks.reqif").write_text(xml)
        record = self.workspace.read_file("blocks.reqif")["records"][0]
        self.assertEqual(record["values"]["AD"], "The unit shall stop and report a fault. shutdown")
        self.assertIn("<p>", record["raw_xhtml"]["AD"])

    def test_reqif_limit_also_bounds_relation_and_hierarchy_collections(self):
        for tag in (
            "SPEC-RELATION",
            "SPEC-HIERARCHY",
            "SPEC-OBJECT-TYPE",
            "DATATYPE-DEFINITION-STRING",
        ):
            with self.subTest(tag=tag):
                (self.root / "many.reqif").write_text(
                    "<REQ-IF>"
                    + "".join(f'<{tag} IDENTIFIER="{i}"/>' for i in range(3))
                    + "</REQ-IF>"
                )
                with patch("arc_engineering.files.MAX_RECORDS", 2), self.assertRaises(ValueError):
                    self.workspace.read_file("many.reqif")

    def test_formula_warning_volume_does_not_break_small_read_pages(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["ID", "Text", "Limit"])
        for index in range(30):
            sheet.append([str(index), "Requirement", "=1+2"])
        workbook.save(self.root / "formulas.xlsx")
        result = self.workspace.read_file("formulas.xlsx", limit=1)
        self.assertEqual(result["total_records"], 30)
        self.assertEqual(len(result["records"]), 1)
        self.assertLessEqual(len(result["warnings"]), 21)
        self.assertTrue(any("30" in warning and "20" in warning for warning in result["warnings"]))

    def test_workbook_stale_dimension_metadata_cannot_hide_records(self):
        workbook = Workbook()
        workbook.active.append(["ID", "Text"])
        workbook.active.append(["R1", "First requirement"])
        workbook.active.append(["R2", "Second requirement"])
        original = self.root / "original.xlsx"
        workbook.save(original)
        with (
            zipfile.ZipFile(original) as source,
            zipfile.ZipFile(self.root / "stale.xlsx", "w") as target,
        ):
            for entry in source.infolist():
                data = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b'dimension ref="A1:B3"', b'dimension ref="A1:A1"')
                target.writestr(entry, data)
        inspection = self.workspace.inspect_file("stale.xlsx")
        self.assertEqual(inspection["sheets"][0]["rows"], 3)
        self.assertEqual(inspection["sheets"][0]["columns"], 2)
        result = self.workspace.import_requirements("stale.xlsx", {"id": "ID", "statement": "Text"})
        self.assertEqual(result["counts"]["imported"], 2)

    def test_saved_import_proposal_has_paged_records_and_empty_table_rejects_cleanly(self):
        (self.root / "req.csv").write_text("ID,Text\nR1,First\nR2,Second\n")
        result = self.workspace.import_requirements("req.csv", {"id": "ID", "statement": "Text"})
        self.workspace.export_artifact("proposal.json", result)
        page = self.workspace.read_file("proposal.json", limit=1)
        self.assertEqual(page["total_records"], 2)
        self.assertEqual(page["records"][0]["values"]["import_id"], "R1")
        self.assertEqual(page["next_offset"], 1)
        for filename in ("empty.csv", "empty.xlsx"):
            with self.subTest(filename=filename), self.assertRaises(ValueError):
                self.workspace.export_artifact(filename, [])
            self.assertFalse((self.root / filename).exists())

    def test_xlsx_export_rejects_truncation_invalid_text_and_renamed_sheets(self):
        for filename, data in [
            ("long.xlsx", [{"text": "x" * 32768}]),
            ("control.xlsx", [{"text": "bad\x00text"}]),
            ("names.xlsx", {"Results": [{"x": 1}], "results": [{"x": 2}]}),
            ("empty.xlsx", {"": [{"x": 1}]}),
        ]:
            with self.subTest(filename=filename), self.assertRaises(ValueError):
                self.workspace.export_artifact(filename, data)
            self.assertFalse((self.root / filename).exists())
        self.workspace.export_artifact("limit.xlsx", [{"text": "x" * 32767}])
        workbook = load_workbook(self.root / "limit.xlsx")
        self.assertEqual(len(workbook.active["A2"].value), 32767)
        workbook.close()


class ModelTests(unittest.TestCase):
    def test_fixed_endpoint_rules_cycles_and_ids(self):
        model = sample_model()
        self.assertTrue(validate_model(model)["valid"])
        model["relationships"][1]["type"] = "verifies"
        self.assertFalse(validate_model(model)["valid"])
        model = sample_model()
        model["relationships"].append(
            {"id": "loop", "type": "derives", "source": "C", "target": "P", "fields": {}}
        )
        self.assertFalse(validate_model(model)["valid"])
        model = sample_model()
        model["items"].append(copy.deepcopy(model["items"][0]))
        self.assertFalse(validate_model(model)["valid"])

    def test_comparison_and_trace_expose_exact_changes_and_recorded_paths(self):
        before = sample_model()
        after = copy.deepcopy(before)
        after["items"][1]["fields"]["statement"] = "Changed"
        difference = compare_models(before, after)
        self.assertEqual(difference["items"]["modified"][0]["id"], "C")
        self.assertEqual(
            difference["items"]["modified"][0]["before"]["fields"]["statement"], "Detail"
        )
        trace = trace_relationships(before, "P", [], "outgoing", 4)
        self.assertTrue(any(path["item_ids"] == ["P", "C", "S"] for path in trace["paths"]))
        self.assertFalse(trace["truncated"])


class CatalogTests(unittest.TestCase):
    def test_catalog_rejects_reference_and_skill_symlinks_outside_package(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "plugin"
            (root / "references").mkdir(parents=True)
            (root / "skills").mkdir()
            outside = Path(directory) / "private.md"
            outside.write_text("Private data")
            link = root / "references/secret.md"
            link.symlink_to(outside)
            with self.assertRaises(ValueError):
                Catalog(root)
            link.unlink()
            outside_skill = Path(directory) / "private-skill"
            outside_skill.mkdir()
            (outside_skill / "SKILL.md").write_text(
                "---\nname: private-skill\ndescription: Private\n---\nPrivate data"
            )
            (root / "skills/private-skill").symlink_to(outside_skill, target_is_directory=True)
            with self.assertRaises(ValueError):
                Catalog(root)

    def test_common_engineering_terms_find_the_specific_workflow(self):
        catalog = Catalog()
        for query, skill_id in [
            ("ConOps", "write-concept-of-operations"),
            ("ICD", "write-interface-control-document"),
            ("RTM", "build-requirement-traceability-matrix"),
            ("functional allocation", "allocate-system-functions"),
            ("functional decomposition", "decompose-system-into-subsystems"),
            ("ECSS requirements", "review-requirement-ecss"),
        ]:
            with self.subTest(query=query):
                self.assertIn(skill_id, {row["id"] for row in catalog.search(query, limit=5)})

    def test_packaged_skill_links_include_references_in_prompts_and_scan_bundle(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "skills/example/references").mkdir(parents=True)
            (root / "references").mkdir()
            (root / "references/engineering-contract.md").write_text("Required guidance")
            (root / "skills/example/references/engineering-contract.md").write_text(
                "Required guidance"
            )
            (root / "skills/example/SKILL.md").write_text(
                "---\nname: example\ndescription: Example task\n---\n"
                "Read [contract](references/engineering-contract.md)."
            )
            catalog = Catalog(root)
            self.assertEqual(catalog.get("example")["references"], ["engineering-contract"])
            self.assertIn("Required guidance", catalog.render("example", "Input"))
            self.assertEqual(len(catalog.scan_entries()[0]["resources"]), 2)

    def test_catalog_is_discoverable_and_extension_digests_match_served_bytes(self):
        catalog = Catalog()
        self.assertEqual(len(catalog.skills), 100)
        self.assertTrue(catalog.search("ECSS individual requirement", limit=5))
        entries = catalog.scan_entries()
        self.assertEqual(len(entries), 5)
        for entry in entries:
            for resource in entry["resources"]:
                text = catalog.read_uri(resource["uri"])
                digest = "sha256:" + hashlib.sha256(text.encode()).hexdigest()
                self.assertEqual(digest, resource["digest"])
        with self.assertRaises(ValueError):
            catalog.read_uri("skill://arc-engineering/../../secret")


if __name__ == "__main__":
    unittest.main()
